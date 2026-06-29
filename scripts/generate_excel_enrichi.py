#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Génère un fichier Excel enrichi téléchargeable depuis le site, en croisant
le CSV OpenData officiel avec les XML d'amendements pour récupérer le
dispositif complet et l'exposé des motifs.

Le fichier produit data/export_amendements_enrichi.xlsx est régénéré
manuellement via le workflow .github/workflows/enrich.yml (déclenchement
workflow_dispatch). Pas d'exécution automatique car c'est lourd
(1 fetch XML par amendement = ~3 minutes).

Trois feuilles : Tous / Dvp durable / Affaires éco. Mêmes colonnes que
l'export Excel standard + colonnes dispositif et exposé des motifs.
"""

from __future__ import annotations
import csv
import io
import json
import re
import html
import sys
import urllib.request
import urllib.error
import ssl
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime, timezone
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    print("⚠ openpyxl non installé. Installer avec : pip install openpyxl", file=sys.stderr)
    sys.exit(1)


# ----- Configuration -----

ROOT = Path(__file__).resolve().parent.parent
OUTPUT_FILE = ROOT / "data" / "export_amendements_enrichi.xlsx"

# Numéro de dossier législatif sur OpenData (PJL 2632 = 54085)
DOSSIER_LEG_ID = "54085"
DOSSIER_CSV_URL = f"https://data.assemblee-nationale.fr/static/openData/repository/17/dossiers_legislatifs_opendata/{DOSSIER_LEG_ID}/excel.csv"

USER_AGENT = "Mozilla/5.0 (compatible; veille-pjl-2632/1.0)"
TIMEOUT = 15
MAX_WORKERS = 10

# Couleurs (réutilisées de generate_excel.py)
GROUP_COLORS = {
    "RN":   {"bg": "1B3D6A", "fg": "FFFFFF"},
    "EPR":  {"bg": "FBC02D", "fg": "1A1816"},
    "LFI":  {"bg": "B71C1C", "fg": "FFFFFF"},
    "SOC":  {"bg": "EF5350", "fg": "FFFFFF"},
    "DR":   {"bg": "1565C0", "fg": "FFFFFF"},
    "EcoS": {"bg": "2E7D32", "fg": "FFFFFF"},
    "DEM":  {"bg": "FF7043", "fg": "FFFFFF"},
    "HOR":  {"bg": "5C6BC0", "fg": "FFFFFF"},
    "LIOT": {"bg": "8D6E63", "fg": "FFFFFF"},
    "UDR":  {"bg": "424242", "fg": "FFFFFF"},
}
STATE_COLORS = {
    "En traitement":   {"bg": "F4F1E8", "fg": "6E6859"},
    "A discuter":      {"bg": "E3EDF6", "fg": "2C5E8A"},
    "Discuté":         {"bg": "ECE4F4", "fg": "6F4D8C"},
    "Adopté":          {"bg": "D5ECDA", "fg": "1A5D2C", "bold": True},
    "Rejeté":          {"bg": "F7D8D4", "fg": "8E2418", "bold": True},
    "Tombé":           {"bg": "ECE9E0", "fg": "5A5448", "bold": True},
    "Non soutenu":     {"bg": "ECE9E0", "fg": "5A5448"},
    "Retiré":          {"bg": "EBE8DF", "fg": "6E6859"},
    "Irrecevable":     {"bg": "F7E3E0", "fg": "8E2418"},
    "Irrecevable 40":  {"bg": "F7E3E0", "fg": "8E2418"},
}
HEADER_BG = "1D4D2C"
HEADER_FG = "FFFFFF"

# Mapping des codes XML de groupe parlementaire
GROUP_REF_MAP = {
    "PO845401": "RN", "PO845407": "EPR", "PO845413": "LFI",
    "PO845419": "SOC", "PO845425": "DR", "PO845439": "EcoS",
    "PO845454": "DEM", "PO845470": "HOR", "PO845485": "LIOT",
    "PO872880": "UDR",
}

# Namespace XML utilisé par OpenData
XML_NS_PREFIX = "{http://schemas.assemblee-nationale.fr/referentiel}"


# ----- Helpers -----

def safe_text(s) -> str:
    """Retire les caractères de contrôle qui font crasher openpyxl."""
    if s is None:
        return ""
    s = str(s)
    return "".join(c for c in s if c == "\n" or c == "\t" or ord(c) >= 32)


def clean_html(raw_html: str) -> str:
    """Décode entités HTML et retire les balises pour produire un texte propre.

    Adapté du script original de l'utilisateur, simplifié pour ne pas
    dépendre de BeautifulSoup (nous utilisons regex+html.unescape).
    """
    if not raw_html:
        return ""
    # Premier décodage des entités numériques (&#x00E8; -> è)
    text = html.unescape(str(raw_html))
    # Retirer les balises HTML résiduelles
    text = re.sub(r"<[^>]+>", " ", text)
    # Second décodage de sécurité
    text = html.unescape(text)
    # Espaces insécables et normalisation
    text = text.replace("\xa0", " ")
    text = re.sub(r"\s+", " ", text).strip()
    return text


def _ssl_context(verified: bool = True):
    """Contexte SSL : vérifié (certifi si dispo) ou — en repli — non vérifié."""
    if not verified:
        return ssl._create_unverified_context()
    try:
        import certifi
        return ssl.create_default_context(cafile=certifi.where())
    except Exception:
        return ssl.create_default_context()


def _uopen(req, timeout=TIMEOUT):
    """urlopen avec repli SSL non vérifié si l'AN sert un certificat non vérifiable
    (WAF/anti-bot intermittent). Donnée PUBLIQUE en lecture seule.
    NB : urlopen ENVELOPPE l'erreur SSL dans URLError -> on teste e.reason."""
    try:
        return urllib.request.urlopen(req, timeout=timeout, context=_ssl_context(True))
    except urllib.error.URLError as e:
        if isinstance(getattr(e, "reason", None), ssl.SSLCertVerificationError):
            print("  certificat AN non vérifiable — repli SSL non vérifié (donnée publique)",
                  file=sys.stderr)
            return urllib.request.urlopen(req, timeout=timeout, context=_ssl_context(False))
        raise


def http_get(url: str, decode: str | None = None) -> bytes | str | None:
    """GET HTTP avec User-Agent et timeout. Retourne None en cas d'échec."""
    try:
        req = urllib.request.Request(url, headers={"User-Agent": USER_AGENT})
        with _uopen(req, timeout=TIMEOUT) as resp:
            raw = resp.read()
            return raw.decode(decode) if decode else raw
    except Exception as e:
        print(f"  GET {url} → {e}", file=sys.stderr)
        return None


# ----- Téléchargement et parsing XML -----

def fetch_amendment_xml(xml_url: str) -> dict:
    """Fetch un XML d'amendement et extrait dispositif + exposé des motifs.

    Renvoie un dict avec les clés : dispositif, expose, tri (vide si erreur).
    """
    raw = http_get(xml_url)
    if not raw:
        return {"dispositif": "", "expose": "", "tri": ""}

    import xml.etree.ElementTree as ET
    try:
        root = ET.fromstring(raw)
    except ET.ParseError:
        return {"dispositif": "", "expose": "", "tri": ""}

    out = {"dispositif": "", "expose": "", "tri": ""}

    disp = root.find(f".//{XML_NS_PREFIX}dispositif")
    if disp is not None and disp.text:
        out["dispositif"] = clean_html(disp.text)

    expose = root.find(f".//{XML_NS_PREFIX}exposeSommaire")
    if expose is not None and expose.text:
        out["expose"] = clean_html(expose.text)

    tri = root.find(f".//{XML_NS_PREFIX}triAmendement")
    if tri is not None and tri.text:
        out["tri"] = tri.text.strip()

    return out


# ----- Mise en forme Excel -----

def write_sheet(ws, rows: list[dict], full_title: str):
    """Écrit une feuille avec les amendements enrichis.

    rows : liste de dicts du CSV avec en plus dispositif/expose/tri
    """
    # Titre
    ws["A1"] = f"Bulletin de veille — PJL n° 2632 — Souveraineté agricole — {full_title}"
    ws["A1"].font = Font(name="Calibri", size=14, bold=True, color="1A1816")
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
    ws.merge_cells("A1:I1")
    ws.row_dimensions[1].height = 24

    ws["A2"] = f"Export OpenData enrichi — généré le {datetime.now(timezone.utc).strftime('%d/%m/%Y à %H:%M UTC')} — {len(rows)} amendements"
    ws["A2"].font = Font(name="Calibri", size=10, italic=True, color="6E6859")
    ws.merge_cells("A2:I2")
    ws.row_dimensions[2].height = 18

    # En-têtes ligne 4
    headers = [
        "N°", "Article", "Auteur", "Groupe", "État", "Tri",
        "Dispositif", "Exposé des motifs", "Lien officiel"
    ]
    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=4, column=col_idx, value=h)
        cell.font = Font(name="Calibri", size=11, bold=True, color=HEADER_FG)
        cell.fill = PatternFill("solid", fgColor=HEADER_BG)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[4].height = 28

    # Bordure légère
    thin = Side(border_style="thin", color="DDDDDD")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    # Tri primaire par triAmendement (= clé alphabétique fournie par l'AN
    # pour ordonner les amendements dans l'ordre de discussion en commission).
    # Exemple : "aaaaaaaaaa" vient avant "aaaaaaaab" qui vient avant "ab".
    # Les amendements sans tri assigné (très récents) finissent en queue,
    # classés par préfixe + numéro.
    def sort_key(r):
        tri = (r.get("tri") or "").strip().lower()
        num = r.get("num", "")
        prefix = "".join(c for c in num if c.isalpha())
        digits = "".join(c for c in num if c.isdigit())
        if tri:
            # Groupe 0 : amendements avec tri assigné, classés alphabétiquement
            return (0, tri, prefix, int(digits) if digits.isdigit() else 0)
        # Groupe 1 : amendements sans tri, classés par préfixe + numéro
        return (1, "", prefix, int(digits) if digits.isdigit() else 0)

    sorted_rows = sorted(rows, key=sort_key)

    for row_idx, r in enumerate(sorted_rows, start=5):
        # N° (avec hyperlien)
        cell = ws.cell(row=row_idx, column=1, value=safe_text(r.get("num", "")))
        cell.font = Font(name="Calibri", size=10, bold=True, color="1565C0", underline="single")
        cell.alignment = Alignment(horizontal="left", vertical="top")
        url = r.get("url", "")
        if url:
            cell.hyperlink = url
        cell.border = border

        # Article
        cell = ws.cell(row=row_idx, column=2, value=safe_text(r.get("article", "")))
        cell.font = Font(name="Calibri", size=10)
        cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        cell.border = border

        # Auteur
        author = r.get("author", "")
        if r.get("rapporteur"):
            author += " (rapporteur)"
        cell = ws.cell(row=row_idx, column=3, value=safe_text(author))
        cell.font = Font(name="Calibri", size=10)
        cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        cell.border = border

        # Groupe (avec coloration)
        group = r.get("group", "")
        cell = ws.cell(row=row_idx, column=4, value=safe_text(group))
        gcol = GROUP_COLORS.get(group, {"bg": "E0E0E0", "fg": "1A1816"})
        cell.font = Font(name="Calibri", size=10, bold=True, color=gcol["fg"])
        cell.fill = PatternFill("solid", fgColor=gcol["bg"])
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border

        # État (avec coloration)
        state = r.get("state", "En traitement")
        cell = ws.cell(row=row_idx, column=5, value=safe_text(state))
        scol = STATE_COLORS.get(state, {"bg": "F4F1E8", "fg": "6E6859"})
        cell.font = Font(name="Calibri", size=10, color=scol["fg"], bold=scol.get("bold", False))
        cell.fill = PatternFill("solid", fgColor=scol["bg"])
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border

        # Tri
        cell = ws.cell(row=row_idx, column=6, value=safe_text(r.get("tri", "")))
        cell.font = Font(name="Calibri", size=10)
        cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        cell.border = border

        # Dispositif (long)
        cell = ws.cell(row=row_idx, column=7, value=safe_text(r.get("dispositif", "")))
        cell.font = Font(name="Calibri", size=10)
        cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        cell.border = border

        # Exposé des motifs (long)
        cell = ws.cell(row=row_idx, column=8, value=safe_text(r.get("expose", "")))
        cell.font = Font(name="Calibri", size=10)
        cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        cell.border = border

        # Lien officiel
        cell = ws.cell(row=row_idx, column=9, value="Voir sur l'AN" if url else "")
        if url:
            cell.font = Font(name="Calibri", size=10, color="1565C0", underline="single")
            cell.hyperlink = url
        else:
            cell.font = Font(name="Calibri", size=10)
        cell.alignment = Alignment(horizontal="left", vertical="top")
        cell.border = border

    # Largeurs de colonnes
    widths = [10, 20, 24, 8, 16, 12, 70, 70, 16]
    for col_idx, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = w

    for row in ws.iter_rows(min_row=5, max_row=ws.max_row):
        ws.row_dimensions[row[0].row].height = 80

    # Filtres + figeage
    last_col = get_column_letter(len(headers))
    last_row = 4 + len(sorted_rows)
    ws.auto_filter.ref = f"A4:{last_col}{last_row}"
    ws.freeze_panes = "A5"


# ----- Pipeline principal -----

def main():
    print(f"=== Génération de l'export OpenData enrichi à {datetime.now(timezone.utc).isoformat()} ===")

    # 1. Télécharger le CSV OpenData
    print(f"\n[1/4] Téléchargement du CSV OpenData (dossier {DOSSIER_LEG_ID})...")
    raw = http_get(DOSSIER_CSV_URL)
    if not raw:
        print("❌ CSV OpenData injoignable. Abandon.")
        sys.exit(1)
    try:
        text = raw.decode("cp1252")
    except UnicodeDecodeError:
        text = raw.decode("utf-8", errors="replace")
    csv_rows = list(csv.DictReader(io.StringIO(text), delimiter=";"))
    print(f"   ✓ {len(csv_rows)} amendements dans le CSV officiel")

    # 2. Pour chaque amendement, fetcher son XML et extraire dispositif + exposé
    print(f"\n[2/4] Téléchargement et parsing des {len(csv_rows)} XML en parallèle...")

    def worker(idx_row):
        idx, row = idx_row
        xml_url = (row.get("URL Amendement format XML") or "").strip()
        if not xml_url:
            return idx, {"dispositif": "", "expose": "", "tri": ""}
        return idx, fetch_amendment_xml(xml_url)

    enriched_xml = {}
    completed = 0
    with ThreadPoolExecutor(max_workers=MAX_WORKERS) as ex:
        futures = {ex.submit(worker, (i, r)): i for i, r in enumerate(csv_rows)}
        for fut in as_completed(futures):
            idx, parsed = fut.result()
            enriched_xml[idx] = parsed
            completed += 1
            if completed % 200 == 0:
                print(f"   ... {completed}/{len(csv_rows)}")
    print(f"   ✓ {len(enriched_xml)} XML traités")

    # 3. Charger amendments.json pour les groupes parlementaires (déjà résolus)
    amendments_json_file = ROOT / "data" / "amendments.json"
    group_by_num = {}
    if amendments_json_file.exists():
        with amendments_json_file.open(encoding="utf-8") as f:
            data = json.load(f)
        for a in data.get("amendments", []):
            group_by_num[a["num"]] = a.get("group", "")

    # 4. Construire la structure pour le rendu Excel
    print(f"\n[3/4] Préparation des données...")
    rows_all = []
    for idx, csv_row in enumerate(csv_rows):
        num = (csv_row.get("Numéro de l'amendement") or "").strip()
        if not num:
            continue
        author_raw = (csv_row.get("Auteur") or "").strip()
        is_rapporteur = bool(re.search(r"rapporteur", author_raw, re.IGNORECASE))
        author = re.sub(
            r"\s+rapporteur(e)?\s+(pour\s+avis\s+)?(au\s+nom\s+de\s+la\s+commission.*)?$",
            "", author_raw, flags=re.IGNORECASE
        ).strip()
        author = re.sub(r"\s+rapporteur(e)?\s*$", "", author, flags=re.IGNORECASE).strip()

        sort = (csv_row.get("Sort de l'amendement") or "").strip()
        state = sort if sort and sort != "Non renseigné" else "En traitement"

        xml_data = enriched_xml.get(idx, {"dispositif": "", "expose": "", "tri": ""})

        rows_all.append({
            "num": num,
            "article": (csv_row.get("Désignation de l'article") or "À classer").strip(),
            "author": author,
            "rapporteur": is_rapporteur,
            "group": group_by_num.get(num, ""),
            "state": state,
            "instance": (csv_row.get("Instance") or "").strip(),
            "url": (csv_row.get("URL Amendement") or "").strip(),
            "tri": xml_data["tri"],
            "dispositif": xml_data["dispositif"],
            "expose": xml_data["expose"],
        })

    # 5. Écrire le fichier Excel à 3 feuilles
    print(f"\n[4/4] Génération du fichier Excel...")
    cd_rows = [r for r in rows_all if r["instance"] == "Développement durable"]
    ce_rows = [r for r in rows_all if r["instance"] == "Affaires économiques"]

    wb = Workbook()
    ws_all = wb.active
    ws_all.title = "Tous"
    write_sheet(ws_all, rows_all, "Tous les amendements")

    if cd_rows:
        ws_cd = wb.create_sheet("Dvp durable")
        write_sheet(ws_cd, cd_rows, "Commission Développement durable")
    if ce_rows:
        ws_ce = wb.create_sheet("Affaires éco")
        write_sheet(ws_ce, ce_rows, "Commission Affaires économiques")

    wb.save(OUTPUT_FILE)

    import os
    size_kb = os.path.getsize(OUTPUT_FILE) / 1024
    print(f"\n🎉 Fichier généré : {OUTPUT_FILE.relative_to(ROOT)} ({size_kb:.0f} ko)")
    print(f"   - Feuille 1 : Tous ({len(rows_all)} amendements)")
    print(f"   - Feuille 2 : Dvp durable ({len(cd_rows)})")
    print(f"   - Feuille 3 : Affaires éco ({len(ce_rows)})")


if __name__ == "__main__":
    main()
